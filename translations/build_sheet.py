from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

NL = "\n"
# (key, page, section, type, english, note, source)
R = []
def add(key, page, section, typ, en, note="", src=""):
    R.append((key, page, section, typ, en, note, src))

# ── Shared: navigation, footer, accessibility
S = "Every page"
add("nav.home", S, "Menu", "Link", "Home", "Top menu and mobile menu.", "Navbar.jsx")
add("nav.about", S, "Menu", "Link", "About Us", "Top menu, mobile menu, footer links, and the Home page 'About Us' button.", "Navbar.jsx, Footer.jsx, Home.jsx")
add("nav.products", S, "Menu", "Link", "Our Products", "Top menu, mobile menu, footer links.", "Navbar.jsx, Footer.jsx")
add("nav.impact", S, "Menu", "Link", "Our Impact", "Top menu, mobile menu, footer links.", "Navbar.jsx, Footer.jsx")
add("nav.partner", S, "Menu", "Button", "Partner", "Short button in the top menu; opens an email. Keep it to one or two words.", "Navbar.jsx")
add("footer.cta_title", S, "Footer", "Heading", "Partner with <i>purpose.</i>", "Large footer headline. The words inside <i>…</i> are shown in a different colour; keep the tags around the matching word(s).", "Footer.jsx")
add("footer.cta_button", S, "Footer", "Button", "Get in touch", "Opens an email to info@goldenkitchengarden.com.", "Footer.jsx")
add("footer.org_heading", S, "Footer", "Heading", "Organization", "Footer column title.", "Footer.jsx")
add("footer.address", S, "Footer", "Label", "Nkotsi, Musanze, Rwanda", "Place names normally stay as they are.", "Footer.jsx")
add("footer.rdb_label", S, "Footer", "Label", "RDB Code: {code}", "Also on the About page. Keep {code}; the number 112368548 is filled in automatically.", "Footer.jsx, About.jsx")
add("footer.links_heading", S, "Footer", "Heading", "Quick Links", "Footer column title.", "Footer.jsx")
add("footer.leadership_heading", S, "Footer", "Heading", "Leadership", "Footer column title.", "Footer.jsx")
add("footer.founder_role", S, "Footer", "Label", "Managing Director & Founder", "Job title shown under Jean de Dieu TWAGIRIMANA.", "Footer.jsx, seo.js")
add("footer.contact_heading", S, "Footer", "Heading", "Contact Us", "Footer column title.", "Footer.jsx")
add("footer.whatsapp", S, "Footer", "Link", "WhatsApp: {phone}", "Keep {phone}; the number is filled in automatically.", "Footer.jsx")
add("footer.phone", S, "Footer", "Link", "Phone: {phone}", "Keep {phone}; the number is filled in automatically.", "Footer.jsx")
add("footer.copyright", S, "Footer", "Label", "© {year} Golden Kitchen Garden Rwanda. All Rights Reserved.", "Keep {year}; the current year is filled in automatically.", "Footer.jsx")
add("a11y.skip_link", S, "Accessibility", "Screen reader", "Skip to content", "Only appears for keyboard and screen-reader users.", "App.jsx")
add("a11y.menu_toggle", S, "Accessibility", "Screen reader", "Toggle menu", "Read aloud for the mobile menu button (☰).", "Navbar.jsx")
add("a11y.main_nav", S, "Accessibility", "Screen reader", "Main", "Read aloud as the name of the main navigation.", "Navbar.jsx")
add("a11y.logo_alt", S, "Accessibility", "Image description", "GKG Logo", "Description of the logo image.", "Navbar.jsx")
add("lang.switch_label", S, "Language switcher (new)", "Screen reader", "Language", "NEW: label for the language switcher that will be added.", "(new)")
add("lang.english", S, "Language switcher (new)", "Button", "English", "NEW: name of the English option in the switcher.", "(new)")
add("lang.kinyarwanda", S, "Language switcher (new)", "Button", "Kinyarwanda", "NEW: name of the Kinyarwanda option in the switcher.", "(new)")

# ── Home
P = "Home"
add("home.hero.eyebrow", P, "Hero (top banner)", "Label", "Est. 2020 · Musanze, Rwanda", "Small line above the main headline. 'Est.' means 'Established'.", "Home.jsx:102")
add("home.hero.title_line1", P, "Hero (top banner)", "Heading", "Empowering Future Generations", "Main headline, line 1. Line 2 is the next row; the two are read together.", "Home.jsx:104")
add("home.hero.title_line2", P, "Hero (top banner)", "Heading", "through Regenerative & Climate-Smart Agriculture", "Main headline, line 2 (softer style).", "Home.jsx:105")
add("home.hero.subtitle", P, "Hero (top banner)", "Paragraph", "Transforming Rwanda's urban and rural spaces into resilient, productive, and beautiful edible landscapes for communities, investors, and the planet.", "", "Home.jsx:108")
add("home.hero.cta_explore", P, "Hero (top banner)", "Button", "Explore Initiatives", "Scrolls down to the programs.", "Home.jsx:110")
add("home.programs.heading", P, "Core programs", "Heading", "Our Core Programs", "", "Home.jsx:122")
add("home.programs.csa.title", P, "Core programs", "Heading", "Regenerative and" + NL + "Climate-Smart Agriculture", "The line break sits where the cell breaks the line; put yours where it reads naturally.", "Home.jsx:130")
add("home.programs.csa.text", P, "Core programs", "Paragraph", "We advance climate resilience by embedding CSA principles into every community-led initiative — using no-tillage techniques, water-smart irrigation, and biodiversity planting to minimise environmental footprint while maximizing yield per square meter.", "CSA = Climate-Smart Agriculture.", "Home.jsx:131")
add("home.programs.nutrition.title", P, "Core programs", "Heading", "Nutrition &" + NL + "Food Security", "Line break optional.", "Home.jsx:138")
add("home.programs.nutrition.text", P, "Core programs", "Paragraph", "We empower smallholder farmers through practical Farmer Field Schools, providing hands-on training from nursery establishment to sustainable crop production. The program promotes nutrition, food security, and climate-smart agriculture while engaging youth, women, and persons with disabilities through secondary school clubs, VSLAs, and cooperatives.", "VSLA = Village Savings and Loan Association; see the Glossary tab.", "Home.jsx:139")
add("home.programs.innovation.title", P, "Core programs", "Heading", "Agrifood" + NL + "Innovation", "Line break optional.", "Home.jsx:143")
add("home.programs.innovation.text", P, "Core programs", "Paragraph", "Integrating IoT-based drip irrigation, remote crop monitoring, and precision composting to transform traditional agro-ecosystems into data-driven productive units.", "IoT = internet-connected sensors; the term can stay in English.", "Home.jsx:144")
add("home.programs.circular.title", P, "Core programs", "Heading", "Circular" + NL + "Economy", "Line break optional.", "Home.jsx:148")
add("home.programs.circular.text", P, "Core programs", "Paragraph", "From kitchen waste to premium compost, we close the nutrient loop. Our reuse/recycle model cuts input costs by up to 60% while regenerating soil health.", "", "Home.jsx:149")
add("home.programs.landscaping.title", P, "Core programs", "Heading", "Commercial Landscaping", "", "Home.jsx:153")
add("home.programs.landscaping.text", P, "Core programs", "Paragraph", "High-end, edible landscaping for private estates, luxury hotels, and institutions — promoting our \"beauty-meets-nutrition\" philosophy where every garden feeds and inspires.", "\"beauty-meets-nutrition\" is the company motto; match your translation of the Glossary entry.", "Home.jsx:154")
add("home.programs.learn_more", P, "Core programs", "Button", "Learn More →", "Keep the arrow.", "Home.jsx:155")
add("home.services.kicker", P, "Premium services", "Label", "Premium Services", "Small label above the headline.", "Home.jsx:166")
add("home.services.title", P, "Premium services", "Heading", "Where Beauty" + NL + "Meets Nutrition.", "Company motto; see the Glossary tab.", "Home.jsx:167")
add("home.services.text", P, "Premium services", "Paragraph", "We design and construct breathtaking edible landscapes that are not only visually stunning but abundantly productive. Every leaf, every pathway, every raised bed is crafted with intention.", "", "Home.jsx:169")
add("home.services.design.title", P, "Premium services", "Heading", "Edible Garden Design & Installation", "", "Home.jsx:174")
add("home.services.design.text", P, "Premium services", "Label", "For private homes, estates, hotels & restaurants", "", "Home.jsx:175")
add("home.services.seedlings.title", P, "Premium services", "Heading", "Organic Vegetable Seedlings", "", "Home.jsx:180")
add("home.services.seedlings.text", P, "Premium services", "Label", "Certified chemical-free, grown in our nursery", "", "Home.jsx:181")
add("home.services.compost.title", P, "Premium services", "Heading", "Premium Organic Compost", "", "Home.jsx:186")
add("home.services.compost.text", P, "Premium services", "Label", "High-quality soil amendment for commercial farms", "", "Home.jsx:187")
add("home.services.consultancy.title", P, "Premium services", "Heading", "Agricultural Consultancy", "", "Home.jsx:192")
add("home.services.consultancy.text", P, "Premium services", "Label", "Training, planning & field support", "", "Home.jsx:193")
add("home.services.cta", P, "Premium services", "Button", "Request a Consultation", "Opens an email.", "Home.jsx:197")
add("home.partners.kicker", P, "Partners", "Label", "Our Ecosystem", "Small label above the headline.", "Home.jsx:211")
add("home.partners.heading", P, "Partners", "Heading", "Partners & Collaborators", "", "Home.jsx:212")
add("home.partners.lede", P, "Partners", "Paragraph", "We work alongside a growing network of government agencies, international organizations, and private sector partners to scale impact across Rwanda.", "", "Home.jsx:215")
add("home.partners.government.title", P, "Partners", "Heading", "Government", "", "Home.jsx:222")
add("home.partners.government.text", P, "Partners", "Label", "Rwanda Agriculture Board (RAB), MINAGRI, Local Government, District Authorities", "Use the official Kinyarwanda names where they exist. Keep the abbreviations RAB and MINAGRI.", "Home.jsx:223")
add("home.partners.ngo.title", P, "Partners", "Heading", "International NGOs", "", "Home.jsx:227")
add("home.partners.ngo.text", P, "Partners", "Label", "UN Agencies, Development Partners, and International Development Organizations", "", "Home.jsx:228")
add("home.partners.private.title", P, "Partners", "Heading", "Private Sector", "", "Home.jsx:232")
add("home.partners.private.text", P, "Partners", "Label", "Luxury Hotels, Restaurants, Private Estates, Commercial Farms & Agribusinesses", "", "Home.jsx:233")
add("home.partners.education.title", P, "Partners", "Heading", "Education", "", "Home.jsx:237")
add("home.partners.education.text", P, "Partners", "Label", "Primary & Secondary Schools, Universities, and Vocational Training Centers across Musanze", "", "Home.jsx:238")
add("home.partners.community.title", P, "Partners", "Heading", "Community Groups", "", "Home.jsx:242")
add("home.partners.community.text", P, "Partners", "Label", "Women's Cooperatives, Youth Associations, and Persons with Disabilities (PWD) Groups", "", "Home.jsx:243")
add("home.partners.technology.title", P, "Partners", "Heading", "Technology", "", "Home.jsx:247")
add("home.partners.technology.text", P, "Partners", "Label", "IoT & AgriTech Providers, Digital Agriculture Platforms, and Research Institutions", "", "Home.jsx:248")
add("home.grow.kicker", P, "What we grow", "Label", "From Our Gardens", "Small label above the headline.", "Home.jsx:259")
add("home.grow.heading", P, "What we grow", "Heading", "What We" + NL + "<i>Grow.</i>", "Displayed on two lines; the words inside <i>…</i> are in italics. Keep the tags.", "Home.jsx:261-262")
add("home.grow.lede", P, "What we grow", "Paragraph", "All of our produce is 100% organic and chemical-free, grown using climate-smart techniques. From leafy greens to companion flowers, every crop serves a purpose in our integrated food systems.", "", "Home.jsx:266")
for i, crop in enumerate(["Kale & Collard Greens", "Onions & Spring Onions", "Bush Beans & Climbing Beans", "Parsley & Dill", "Cabbage & Bok Choy", "Sweet Potato", "Amaranth Greens", "Marigolds (Companion)", "Fennel", "Organic Compost"], 1):
    add(f"home.grow.crop_{i:02d}", P, "What we grow", "Tag", crop, "Crop name shown as a short tag. Use the common local name if farmers use one." if i == 1 else "Crop name tag.", f"Home.jsx:{270+i}")
add("home.vision.kicker", P, "Vision 2030", "Label", "Our Vision 2030", "", "Home.jsx:289")
add("home.vision.title", P, "Vision 2030", "Heading", "Scale.", "One very large word.", "Home.jsx:290")
add("home.vision.stat1", P, "Vision 2030", "Label", "Trained Beneficiaries", "Shown under the figure 12,000+ (the figure is not translated).", "Home.jsx:295")
add("home.vision.stat2", P, "Vision 2030", "Label", "Kitchen Gardens Installed", "Shown under the figure 1,000+.", "Home.jsx:299")
add("home.vision.stat3", P, "Vision 2030", "Label", "Institutional Partnerships", "Shown under the figure 150.", "Home.jsx:303")
add("home.vision.text", P, "Vision 2030", "Paragraph", "By 2030, GKG aims to be Rwanda's <i>National Hub for Urban Agricultural Excellence</i> — scaling from Musanze into every major city, partnering with NGOs, government, and private investors to permanently transform the nation's food landscape.", "The words inside <i>…</i> are in italics. Keep the tags around the matching words.", "Home.jsx:307")
add("home.gallery.kicker", P, "Gallery", "Label", "Journal & Work", "Small label above the headline.", "Home.jsx:317")
add("home.gallery.heading", P, "Gallery", "Heading", "Impact in Action.", "", "Home.jsx:318")
add("home.gallery.lede", P, "Gallery", "Paragraph", "A visual diary of our daily operations, training sessions, and the communities we empower across Rwanda.", "", "Home.jsx:320")

# ── About
P = "About Us"
add("about.kicker", P, "Page header", "Label", "Corporate Overview", "Small label above the page title.", "About.jsx:20")
add("about.title", P, "Page header", "Heading", "About Us", "Page title. Normally the same as the menu item.", "About.jsx:21")
add("about.who.kicker", P, "Who we are", "Label", "Who we are", "", "About.jsx:29")
add("about.who.title", P, "Who we are", "Heading", "Bridging the Gap Between Rapid Urbanization & Food Security.", "", "About.jsx:30")
add("about.who.text1", P, "Who we are", "Paragraph", "Golden Kitchen Garden Rwanda Ltd (GKG) is a registered social enterprise established to promote climate-smart agriculture, food security, nutrition improvement, environmental sustainability, and economic empowerment in Rwanda.", "", "About.jsx:32")
add("about.who.text2", P, "Who we are", "Paragraph", "Founded in 2020 by Jean de Dieu TWAGIRIMANA with the vision of transforming underutilized spaces into productive and sustainable food systems that improve livelihoods and resilience among communities. We operate through innovative extension on climate-smart agricultural approaches including kitchen gardens, school gardens, organic farming, circular economy solutions, digital agriculture, and value chain development.", "", "About.jsx:35")
add("about.who.hq", P, "Who we are", "Label", "HQ: Nkotsi, Musanze", "HQ = headquarters. (The 'RDB Code' chip reuses footer.rdb_label.)", "About.jsx:39")
add("about.vision.title", P, "Vision & mission", "Heading", "Our Vision", "", "About.jsx:53")
add("about.vision.text", P, "Vision & mission", "Paragraph", "A climate-resilient Rwanda free from hunger where every community has access to sustainable and nutritious food systems.", "", "About.jsx:54")
add("about.mission.title", P, "Vision & mission", "Heading", "Our Mission", "", "About.jsx:57")
add("about.mission.text", P, "Vision & mission", "Paragraph", "To improve livelihoods through regenerative climate-smart agriculture, nutrition systems, environmental sustainability, and inclusive value chains driven by innovative farming systems.", "", "About.jsx:58")
add("about.objectives.kicker", P, "Strategic objectives", "Label", "Goals & Structure", "", "About.jsx:68")
add("about.objectives.heading", P, "Strategic objectives", "Heading", "Strategic Objectives", "", "About.jsx:69")
objectives = [
    ("Strengthen nutrition and food security", "Focusing on vulnerable populations through sustainable means."),
    ("Expand climate-smart kitchen gardens", "Deploying models in both households and institutions."),
    ("Promote sustainable agricultural technologies", "Leveraging innovation for resilient food systems."),
    ("Increase income generation opportunities", "Focusing specifically on empowering women, youth, and persons with disabilities."),
    ("Develop strong market linkages", "Enhancing value chains for our beneficiaries."),
    ("Enhance institutional capacity", "Ensuring the sustainability and growth of GKG operations."),
    ("Scale operations and Job Creation", "Aiming to create up to 12,000 jobs by 2030."),
]
for i, (t, d) in enumerate(objectives, 1):
    add(f"about.objectives.{i}.title", P, "Strategic objectives", "Heading", t, f"Objective {i}, title.", f"About.jsx:{5+i}")
    add(f"about.objectives.{i}.text", P, "Strategic objectives", "Label", d, f"Objective {i}, description.", f"About.jsx:{5+i}")
add("about.leadership.title", P, "Leadership", "Heading", "Leadership.", "One large word.", "About.jsx:88")
add("about.leadership.text", P, "Leadership", "Paragraph", "GKG adopts a functional organizational structure that promotes efficiency and accountability. Governed by a <b>Board of Directors</b> providing strategic oversight, daily operations are led by the <b>Managing Director</b>, supported by dedicated department heads across Finance, Programs, MEAL, and Procurement.", "The words inside <b>…</b> are bold. Keep the tags around the matching words. MEAL = Monitoring, Evaluation, Accountability & Learning; the abbreviation can stay.", "About.jsx:90")
add("about.leadership.cta", P, "Leadership", "Button", "View Our Products", "", "About.jsx:93")

# ── Products
P = "Our Products"
add("products.kicker", P, "Page header", "Label", "What We Offer", "Small label above the page title.", "Governance.jsx:23")
add("products.title", P, "Page header", "Heading", "Our" + NL + "Products", "Page title on two lines. Normally matches the 'Our Products' menu item.", "Governance.jsx:24")
add("products.section.kicker", P, "Product list", "Label", "Products & Services", "", "Governance.jsx:32")
add("products.section.heading", P, "Product list", "Heading", "Cultivated for Homes, Farms & Institutions", "", "Governance.jsx:33")
prods = ["Kitchen Garden Design and Installation", "Agroforestry Nursery", "Landscaping", "French Beans Cultivation", "Strawberries Cultivation", "Vegetables and Spices Nursery Seedbeds", "Upcycling", "Maize Production", "Permaculture Design", "Agriculture Consultation", "Agriculture Capacity Building"]
for i, p in enumerate(prods, 1):
    add(f"products.item_{i:02d}", P, "Product list", "Heading", p, "Product card title; also listed in search-engine data." if i == 1 else "Product card title.", f"Governance.jsx:{4+i}, seo.js:{11+i}")
add("products.image_alt", P, "Product list", "Image description", "{product} by Golden Kitchen Garden Rwanda", "Description of each product photo. Keep {product}; the product name from the rows above is filled in.", "Governance.jsx:41")

# ── Impact
P = "Our Impact"
add("impact.kicker", P, "Page header", "Label", "Measured Progress", "Small label above the page title.", "Safeguarding.jsx:9")
add("impact.title", P, "Page header", "Heading", "Our" + NL + "Impact", "Page title on two lines. Normally matches the 'Our Impact' menu item.", "Safeguarding.jsx:10")
add("impact.intro.kicker", P, "Introduction", "Label", "Across Rwanda", "", "Safeguarding.jsx:18")
add("impact.intro.title", P, "Introduction", "Heading", "Growing resilient communities.", "", "Safeguarding.jsx:19")
add("impact.intro.text", P, "Introduction", "Paragraph", "Golden Kitchen Garden Rwanda strengthens livelihoods, food security, and climate resilience through practical agriculture programs, modern kitchen gardens, community organizations, and school-based learning.", "", "Safeguarding.jsx:21")
add("impact.people.title", P, "Impact figures", "Heading", "Women, Youth and Persons with Disabilities", "Card above the figure 3,300+.", "Safeguarding.jsx:35")
add("impact.people.text", P, "Impact figures", "Label", "Empowered through Regenerative & Climate-Smart Agriculture", "", "Safeguarding.jsx:37")
add("impact.gardens.title", P, "Impact figures", "Heading", "Kitchen Gardens", "Card above the figure 400+.", "Safeguarding.jsx:40")
add("impact.gardens.text", P, "Impact figures", "Label", "Modern Kitchen Gardens Installed Across Rwanda", "", "Safeguarding.jsx:42")
add("impact.orgs.title", P, "Impact figures", "Heading", "Community Organizations", "", "Safeguarding.jsx:45")
add("impact.orgs.vslas", P, "Impact figures", "Label", "{n} VSLAs", "Keep {n}; the figure 50+ is filled in.", "Safeguarding.jsx:46")
add("impact.orgs.coops", P, "Impact figures", "Label", "{n} Cooperatives", "Keep {n}; the figure 10 is filled in.", "Safeguarding.jsx:46")
add("impact.orgs.companies", P, "Impact figures", "Label", "{n} Companies", "Keep {n}; the figure 11+ is filled in.", "Safeguarding.jsx:46")
add("impact.orgs.text", P, "Impact figures", "Label", "Supported and served", "", "Safeguarding.jsx:47")
add("impact.schools.title", P, "Impact figures", "Heading", "School Programs", "Card above the figure 30+.", "Safeguarding.jsx:50")
add("impact.schools.text", P, "Impact figures", "Label", "Primary and Secondary School Agriculture Clubs Established", "", "Safeguarding.jsx:52")

# ── 404
P = "Page not found"
add("notfound.title", P, "Error page", "Heading", "Page not found", "Shown when a link is broken.", "NotFound.jsx:9")
add("notfound.text", P, "Error page", "Paragraph", "The page you were looking for doesn't exist.", "", "NotFound.jsx:10")
add("notfound.cta", P, "Error page", "Button", "Back to home", "", "NotFound.jsx:11")

# ── Image descriptions (alt text)
P = "Image descriptions"
alts = [
    ("home.hero", "Golden Kitchen Garden Rwanda farm team tending a strawberry field", "Home.jsx:96, seo.js:222", "Home top photo; also used when the site is shared on social media."),
    ("home.programs.csa", "Regenerative and climate-smart agriculture", "Home.jsx:127", ""),
    ("home.programs.nutrition", "Nutrition and food security programs", "Home.jsx:137", ""),
    ("home.programs.innovation", "Agrifood innovation", "Home.jsx:142", ""),
    ("home.programs.circular", "Circular economy composting", "Home.jsx:147", ""),
    ("home.programs.landscaping", "Commercial edible landscaping", "Home.jsx:152", ""),
    ("home.services", "Edible landscaping pathway designed by Golden Kitchen Garden Rwanda", "Home.jsx:200", ""),
    ("home.gallery_01", "Kitchen garden construction by GKG Rwanda", "Home.jsx:324", ""),
    ("home.gallery_02", "GKG community work in Musanze, Rwanda", "Home.jsx:325", ""),
    ("home.gallery_03", "Community members preparing farmland together with GKG Rwanda", "Home.jsx:326", ""),
    ("home.gallery_04", "Kitchen garden installed by GKG Rwanda", "Home.jsx:327", ""),
    ("home.gallery_05", "Organic vegetable harvest from a GKG garden", "Home.jsx:328", ""),
    ("home.gallery_06", "Women farmers trained by GKG Rwanda", "Home.jsx:329", ""),
    ("home.gallery_07", "Organic vegetable seedlings from the GKG nursery", "Home.jsx:330", ""),
    ("home.gallery_08", "Farmer tending crops in a GKG-supported field", "Home.jsx:331", ""),
    ("home.gallery_09", "Community impact of GKG programs in Rwanda", "Home.jsx:332", ""),
    ("about.who", "Women farmers empowered through GKG climate-smart agriculture in Rwanda", "About.jsx:43", ""),
    ("impact.intro", "Golden Kitchen Garden Rwanda community impact", "Safeguarding.jsx:25", ""),
]
for k, en, src, note in alts:
    add(f"alt.{k}", P, "Photo descriptions", "Image description", en, note or "Read aloud to blind visitors and used by Google Images.", src)

# ── Search engines & sharing
P = "Search & sharing"
seo = [
    ("seo.home.title", "Browser tab title", "Golden Kitchen Garden Rwanda (GKG) | Climate-Smart Agriculture in Musanze", "Home page. Shown in the browser tab and as the Google result title. Aim for under 60 characters."),
    ("seo.home.description", "Google description", "Golden Kitchen Garden Rwanda (GKG) is a Musanze social enterprise building climate-smart kitchen gardens, edible landscapes, organic seedlings and compost, and farmer training for food security across Rwanda.", "Home page. Shown under the title in Google results. Aim for under 160 characters."),
    ("seo.about.title", "Browser tab title", "About Us | Golden Kitchen Garden Rwanda (GKG)", "About page."),
    ("seo.about.description", "Google description", "Founded in 2020 by Jean de Dieu Twagirimana, Golden Kitchen Garden Rwanda Ltd (RDB 112368548) is a registered social enterprise in Nkotsi, Musanze advancing nutrition, food security and climate-smart agriculture.", "About page. Aim for under 160 characters."),
    ("seo.products.title", "Browser tab title", "Products & Services | Kitchen Gardens, Landscaping & Nursery – GKG Rwanda", "Products page."),
    ("seo.products.description", "Google description", "Kitchen garden design and installation, landscaping, permaculture design, agroforestry and vegetable nurseries, French beans, strawberries, maize, upcycling, and agriculture consultation and training in Rwanda.", "Products page. Aim for under 160 characters."),
    ("seo.impact.title", "Browser tab title", "Our Impact | Golden Kitchen Garden Rwanda (GKG)", "Impact page."),
    ("seo.impact.description", "Google description", "3,300+ women, youth and persons with disabilities empowered, 400+ modern kitchen gardens installed, 50+ VSLAs, 10 cooperatives and 30+ school agriculture clubs supported across Rwanda.", "Impact page. Aim for under 160 characters."),
    ("seo.notfound.title", "Browser tab title", "Page not found | Golden Kitchen Garden Rwanda", "Error page."),
    ("seo.notfound.description", "Google description", "This page does not exist. Visit Golden Kitchen Garden Rwanda to learn about our climate-smart agriculture work in Musanze.", "Error page."),
    ("seo.org.description", "Company description", "Golden Kitchen Garden Rwanda (GKG) is a registered social enterprise in Musanze, Rwanda, founded in 2020. It promotes regenerative and climate-smart agriculture, kitchen and school gardens, nutrition and food security, circular-economy composting, edible landscaping, and income opportunities for women, youth, and persons with disabilities.", "Read by Google to describe the company. Visitors don't see it on the page."),
    ("seo.org.slogan", "Company slogan", "Where beauty meets nutrition.", "Company motto as read by Google; match the Glossary entry."),
    ("seo.breadcrumb.home", "Breadcrumb", "Home", "Shown by Google in the page path (Home › About Us). Usually the same as nav.home."),
]
for k, typ, en, note in seo:
    add(k, P, "Google & social media", typ, en, note, "seo.js")

# Glossary: key terms to keep consistent
GLOSS = [
    ("Climate-Smart Agriculture (CSA)", "Farming that adapts to climate change, keeps yields up, and lowers emissions. Appears about 12 times."),
    ("Regenerative agriculture", "Farming that rebuilds soil health instead of depleting it."),
    ("Kitchen garden", "A small vegetable garden next to a home or institution. GKG's core product."),
    ("School garden / school agriculture club", "Gardens and farming clubs run by primary and secondary schools."),
    ("Edible landscaping", "Decorative gardens made of plants you can eat."),
    ("Food security", ""),
    ("Nutrition", ""),
    ("Social enterprise", "A company that exists to solve a social problem."),
    ("Persons with disabilities (PWD)", ""),
    ("Women, youth and persons with disabilities", "Recurring phrase naming the groups GKG serves."),
    ("VSLA (Village Savings and Loan Association)", "Keep the abbreviation if the local term is 'VSLA' too."),
    ("Cooperative", ""),
    ("Farmer Field School", "Hands-on training held in a farmer's field."),
    ("Circular economy", "Reusing waste (e.g. kitchen scraps into compost)."),
    ("Compost / organic compost", ""),
    ("Seedlings / nursery", "Young plants and the place they are raised."),
    ("Agroforestry", "Growing trees together with crops."),
    ("Permaculture", ""),
    ("Upcycling", "Turning waste materials into useful products."),
    ("Value chain", "All the steps from farm to buyer."),
    ("Livelihoods", ""),
    ("Resilience / climate-resilient", ""),
    ("Where beauty meets nutrition", "Company motto. Used in several places; translate once and reuse."),
    ("Partner / partnership", ""),
]
DNT = [
    ("Golden Kitchen Garden Rwanda / Golden Kitchen Garden Rwanda Ltd", "Company name"),
    ("GKG / GKG Rwanda", "Short company name and logo text"),
    ("Jean de Dieu TWAGIRIMANA", "Founder's name"),
    ("Nkotsi, Musanze, Rwanda", "Place names"),
    ("112368548", "RDB company code"),
    ("+250 788 206 976", "Phone and WhatsApp number"),
    ("info@goldenkitchengarden.com / customer@goldenkitchengarden.com", "Email addresses"),
    ("RAB, MINAGRI, MEAL, IoT, CSA", "Abbreviations; you may add the Kinyarwanda full name alongside"),
    ("12,000+ · 1,000+ · 150 · 3,300+ · 400+ · 50+ · 10 · 11+ · 30+ · 60%", "Figures; only the words around them are translated"),
    ("{year} {code} {phone} {product} {n}", "Placeholders that the website fills in; keep them exactly as written"),
]

# ── Build workbook
FONT = "Arial"
navy = "1F4E3D"
hdr_fill = PatternFill("solid", fgColor=navy)
hdr_font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
input_fill = PatternFill("solid", fgColor="FFF6C8")
done_fill = PatternFill("solid", fgColor="DDF0DD")
band_fill = PatternFill("solid", fgColor="F3F6F4")
thin = Side(style="thin", color="C9D3CE")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_top = Alignment(wrap_text=True, vertical="top")
base = Font(name=FONT, size=10)

wb = Workbook()

# Sheet 1: Instructions
ins = wb.active
ins.title = "Instructions"
ins.sheet_view.showGridLines = False
ins.column_dimensions["A"].width = 3
ins.column_dimensions["B"].width = 34
ins.column_dimensions["C"].width = 90
ins["B2"] = "GKG website: Kinyarwanda translation"
ins["B2"].font = Font(name=FONT, bold=True, size=16, color=navy)
ins["B3"] = "goldenkitchengarden.com · every piece of text on the site, one row each"
ins["B3"].font = Font(name=FONT, italic=True, size=10, color="5A6B63")
rows = [
    ("What to fill in", "Only the yellow column 'Kinyarwanda' on the Translations and Glossary tabs. Leave every other column as it is."),
    ("How to start", "Fill in the Glossary tab first, so recurring terms (e.g. 'Climate-Smart Agriculture', 'kitchen garden') are translated the same way everywhere."),
    ("Line breaks", "Some headings break across two lines on the site. Those cells hold two lines; put your line break (Alt+Enter in Excel) where it reads naturally."),
    ("<i>…</i> and <b>…</b>", "Mark words shown in italics or bold. Keep the tags and move them to the matching Kinyarwanda words."),
    ("{year} {code} {phone} {product} {n}", "Placeholders the website fills in automatically. Copy them into your translation unchanged, in whatever position suits the sentence."),
    ("Do not translate", "Names, numbers, phone numbers, and emails stay as they are. See the 'Do not translate' tab."),
    ("Status", "Updates by itself: a row turns green once it has a translation."),
    ("Key column", "The website's reference for each text. Don't edit it; it's how the translations are matched back to the site."),
    ("Unsure about a row?", "Leave it blank or write your question in the 'Translator comment' column."),
]
r = 5
for label, text in rows:
    ins.cell(r, 2, label).font = Font(name=FONT, bold=True, size=10)
    c = ins.cell(r, 3, text); c.font = base; c.alignment = wrap_top
    ins.cell(r, 2).alignment = wrap_top
    r += 1
r += 1
ins.cell(r, 2, "Example of a filled row").font = Font(name=FONT, bold=True, size=11, color=navy); r += 1
for j, h in enumerate(["English", "Kinyarwanda"]):
    c = ins.cell(r, 2 + j, h); c.font = hdr_font; c.fill = hdr_fill
r += 1
ins.cell(r, 2, "Contact Us").font = base
c = ins.cell(r, 3, "Twandikire"); c.font = Font(name=FONT, size=10, italic=True); c.fill = input_fill
ins.cell(r + 1, 2, "(illustration only; your own wording takes priority)").font = Font(name=FONT, size=9, italic=True, color="7A8A83")
r += 3
ins.cell(r, 2, "Progress").font = Font(name=FONT, bold=True, size=11, color=navy); r += 1
total = len(R)
last = total + 1
prog = [
    ("Texts on the site", f"=COUNTA(Translations!F2:F{last})", "0"),
    ("Translated", f"=COUNTA(Translations!G2:G{last})", "0"),
    ("Remaining", f"=C{r}-C{r+1}", "0"),
    ("Done", f"=IF(C{r}=0,0,C{r+1}/C{r})", "0%"),
    ("Glossary terms translated", f"=COUNTA(Glossary!C2:C{len(GLOSS)+1})&\" of \"&COUNTA(Glossary!B2:B{len(GLOSS)+1})", "@"),
]
for label, f, fmt in prog:
    ins.cell(r, 2, label).font = base
    c = ins.cell(r, 3, f); c.font = Font(name=FONT, size=10, bold=True); c.number_format = fmt
    c.alignment = Alignment(horizontal="left")
    r += 1

# Sheet 2: Translations
ws = wb.create_sheet("Translations")
headers = ["#", "Page", "Section", "Type", "English", "Kinyarwanda", "Status", "Notes for translator", "Translator comment", "Key", "Where in code"]
# column letters: A # | B Page | C Section | D Type | E English ... wait: keep F English/G Kinyarwanda for formulas
headers = ["#", "Page", "Section", "Type", "Key", "English", "Kinyarwanda", "Status", "Notes for translator", "Translator comment", "Where in code"]
widths = [5, 16, 20, 15, 30, 55, 55, 12, 45, 30, 26]
for j, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(1, j, h); c.font = hdr_font; c.fill = hdr_fill; c.border = border
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[1].height = 22
prev_section = None
band = False
for i, (key, page, section, typ, en, note, src) in enumerate(R, 2):
    if (page, section) != prev_section:
        band = not band; prev_section = (page, section)
    vals = [i - 1, page, section, typ, key, en, None, f'=IF(LEN(TRIM(G{i}))>0,"Done","To do")', note, None, src]
    for j, v in enumerate(vals, 1):
        c = ws.cell(i, j, v); c.font = base; c.alignment = wrap_top; c.border = border
        if band and j not in (7,): c.fill = band_fill
    ws.cell(i, 7).fill = input_fill
    ws.cell(i, 5).font = Font(name="Consolas", size=9, color="5A6B63")
    ws.cell(i, 11).font = Font(name=FONT, size=8, color="7A8A83")
    ws.cell(i, 6).font = Font(name=FONT, size=10, bold=typ == "Heading")
ws.freeze_panes = "G2"
ws.auto_filter.ref = f"A1:K{last}"
ws.conditional_formatting.add(f"H2:H{last}", FormulaRule(formula=[f'$H2="Done"'], fill=done_fill, font=Font(name=FONT, color="1E6B34", bold=True)))
ws.conditional_formatting.add(f"G2:G{last}", FormulaRule(formula=[f'LEN(TRIM($G2))>0'], fill=done_fill))

# Sheet 3: Glossary
gl = wb.create_sheet("Glossary")
for j, (h, w) in enumerate(zip(["English term", "", "Kinyarwanda", "What it means"], [4, 40, 40, 60]), 1):
    pass
gh = ["#", "English term", "Kinyarwanda", "What it means / where it's used"]
for j, (h, w) in enumerate(zip(gh, [5, 42, 42, 65]), 1):
    c = gl.cell(1, j, h); c.font = hdr_font; c.fill = hdr_fill; c.border = border
    gl.column_dimensions[c.column_letter].width = w
for i, (term, note) in enumerate(GLOSS, 2):
    for j, v in enumerate([i - 1, term, None, note], 1):
        c = gl.cell(i, j, v); c.font = base; c.alignment = wrap_top; c.border = border
    gl.cell(i, 3).fill = input_fill
    gl.cell(i, 2).font = Font(name=FONT, size=10, bold=True)
gl.freeze_panes = "C2"
gl.conditional_formatting.add(f"C2:C{len(GLOSS)+1}", FormulaRule(formula=['LEN(TRIM($C2))>0'], fill=done_fill))

# Sheet 4: Do not translate
dn = wb.create_sheet("Do not translate")
for j, (h, w) in enumerate(zip(["Keep exactly as written", "What it is"], [60, 60]), 1):
    c = dn.cell(1, j, h); c.font = hdr_font; c.fill = hdr_fill; c.border = border
    dn.column_dimensions[c.column_letter].width = w
for i, (t, n) in enumerate(DNT, 2):
    for j, v in enumerate([t, n], 1):
        c = dn.cell(i, j, v); c.font = base; c.alignment = wrap_top; c.border = border

for s in (ws, gl):
    s.sheet_properties.tabColor = "2D6A4F"
wb.active = 1
out = "/home/ec2-user/gkg/translations/gkg-kinyarwanda-translations.xlsx"
wb.save(out)
keys = [r[0] for r in R]
assert len(keys) == len(set(keys)), "duplicate keys"
print(out, len(R), "rows")
